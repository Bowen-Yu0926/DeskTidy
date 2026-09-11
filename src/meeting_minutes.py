"""Create dated meeting-minutes Word documents in a configured folder."""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from src.settings import install_root

MINUTES_DIR_NAME = "纪要"

_SERIAL_RE = re.compile(r"^(\d{8})_(\d+)\.docx$", re.IGNORECASE)


def meeting_minutes_settings(settings: dict) -> dict:
    raw = settings.get("meeting_minutes")
    if not isinstance(raw, dict):
        raw = {}
        settings["meeting_minutes"] = raw
    raw.setdefault("enabled", True)
    raw.setdefault("folder", "")
    return raw


def meeting_minutes_enabled(settings: dict | None) -> bool:
    if not settings:
        return False
    return bool(meeting_minutes_settings(settings).get("enabled", True))


def default_minutes_folder(*, ensure: bool = True) -> Path:
    """Default: <install_root>/纪要 beside the exe (or project root in dev)."""
    folder = install_root() / MINUTES_DIR_NAME
    if ensure:
        folder.mkdir(parents=True, exist_ok=True)
    return folder


def resolve_minutes_folder(settings: dict | None = None, *, ensure: bool = True) -> Path:
    """Return configured minutes folder, or install_root/纪要 when empty."""
    import sys

    folder: Path | None = None
    custom = ""
    if settings:
        cfg = meeting_minutes_settings(settings)
        custom = str(cfg.get("folder") or "").strip()
        if custom:
            folder = Path(custom).expanduser()
    if folder is None:
        folder = default_minutes_folder(ensure=False)
        if not getattr(sys, "frozen", False) and not custom:
            try:
                empty = (not folder.is_dir()) or (not any(folder.iterdir()))
            except OSError:
                empty = True
            if empty:
                alt = install_root() / "dist" / MINUTES_DIR_NAME
                try:
                    if alt.is_dir() and any(alt.iterdir()):
                        folder = alt
                except OSError:
                    pass
    if ensure:
        folder.mkdir(parents=True, exist_ok=True)
    return folder


def get_minutes_folder(settings: dict) -> Path:
    """Resolved minutes directory (custom path or install_root/纪要)."""
    return resolve_minutes_folder(settings, ensure=False)


def open_minutes_folder(settings: dict) -> Path:
    """Open the minutes folder in Explorer; create it if needed."""
    folder = resolve_minutes_folder(settings, ensure=True)
    if not folder.is_dir():
        raise ValueError(f"纪要文件夹不可用：{folder}")
    _open_path(folder)
    return folder


def next_minutes_path(folder: Path, when: datetime | None = None) -> Path:
    """Return ``YYYYMMDD_NNN.docx`` with the next daily serial in ``folder``."""
    day = (when or datetime.now()).strftime("%Y%m%d")
    max_n = 0
    if folder.is_dir():
        for entry in folder.iterdir():
            if not entry.is_file():
                continue
            match = _SERIAL_RE.match(entry.name)
            if not match:
                continue
            if match.group(1) != day:
                continue
            max_n = max(max_n, int(match.group(2)))
    return folder / f"{day}_{max_n + 1:03d}.docx"


def create_meeting_minutes(
    settings: dict,
    *,
    when: datetime | None = None,
    open_after: bool = True,
) -> Path:
    """Create a Word minutes file in the configured folder; return its path."""
    folder = resolve_minutes_folder(settings, ensure=True)
    if not folder.is_dir():
        raise ValueError(f"纪要文件夹不可用：{folder}")

    path = next_minutes_path(folder, when=when)
    now = when or datetime.now()
    _write_minutes_docx(path, now)

    if open_after:
        _open_path(path)
    return path


def _write_minutes_docx(path: Path, when: datetime) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()
    black = RGBColor(0, 0, 0)

    def _apply_yahei(style_name: str, size_pt: float, *, bold: bool | None = None) -> None:
        style = doc.styles[style_name]
        font = style.font
        font.name = "微软雅黑"
        font.size = Pt(size_pt)
        font.color.rgb = black
        if bold is not None:
            font.bold = bold
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:ascii"), "微软雅黑")
        rfonts.set(qn("w:hAnsi"), "微软雅黑")
        rfonts.set(qn("w:eastAsia"), "微软雅黑")
        rfonts.set(qn("w:cs"), "微软雅黑")

    def _format_run(run, *, size_pt: float = 11, bold: bool = False) -> None:
        run.font.name = "微软雅黑"
        run.font.size = Pt(size_pt)
        run.font.bold = bold
        run.font.color.rgb = black
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:ascii"), "微软雅黑")
        rfonts.set(qn("w:hAnsi"), "微软雅黑")
        rfonts.set(qn("w:eastAsia"), "微软雅黑")
        rfonts.set(qn("w:cs"), "微软雅黑")

    _apply_yahei("Normal", 11, bold=False)
    _apply_yahei("Heading 1", 18, bold=True)
    _apply_yahei("Heading 2", 14, bold=True)

    title = doc.add_heading("会议纪要", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        _format_run(run, size_pt=18, bold=True)

    meta = doc.add_paragraph()
    r1 = meta.add_run(f"日期：{when.strftime('%Y年%m月%d日')}")
    _format_run(r1, size_pt=11, bold=True)
    r2 = meta.add_run(f"　　编号：{path.stem}")
    _format_run(r2, size_pt=11, bold=False)

    for heading, hint in (
        ("一、出席人员", "（请填写）"),
        ("二、会议议题", "（请填写）"),
        ("三、讨论内容", "（请填写）"),
        ("四、决议事项", "（请填写）"),
        ("五、待办事项", "（请填写）"),
    ):
        h = doc.add_heading(heading, level=2)
        for run in h.runs:
            _format_run(run, size_pt=14, bold=True)
        p = doc.add_paragraph()
        hr = p.add_run(hint)
        _format_run(hr, size_pt=11, bold=False)

    doc.save(str(path))


def _open_path(path: Path) -> None:
    try:
        import os

        os.startfile(str(path))  # type: ignore[attr-defined]
    except OSError:
        pass
